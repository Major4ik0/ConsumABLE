import os
import re
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime

def make_short_fio(full_name: str) -> str:
    name = full_name.strip()
    m = re.match(r'^([А-ЯЁа-яё\-]+)\s+([А-ЯЁа-яё])\.\s*([А-ЯЁа-яё])\.$', name)
    if m:
        return f"{m.group(1)} {m.group(2).upper()}.{m.group(3).upper()}."
    parts = name.split()
    if len(parts) >= 3:
        i1 = parts[1][0].upper()
        i2 = parts[2][0].upper()
        return f"{parts[0]} {i1}.{i2}."
    elif len(parts) == 2:
        if '.' in parts[1]:
            initials = ''.join([c for c in parts[1] if c.isalpha()]).upper()
            if len(initials) >= 2:
                return f"{parts[0]} {initials[0]}.{initials[1]}."
            elif len(initials) == 1:
                return f"{parts[0]} {initials[0]}."
        return f"{parts[0]} {parts[1][0].upper()}."
    return name

def detect_gender(fio: str, explicit_val=None) -> str:
    if explicit_val is not None:
        v = str(explicit_val).strip().upper()
        if v.startswith("Ж") or v.startswith("F") or "ЖЕН" in v:
            return "Ж"
        if v.startswith("М") or v.startswith("M") or "МУЖ" in v:
            return "М"
    parts = fio.strip().split()
    if len(parts) >= 3:
        patronymic = parts[2].lower()
        if patronymic.endswith("на"):
            return "Ж"
        if patronymic.endswith("ич") or patronymic.endswith("ыч"):
            return "М"
    if parts and (parts[0].lower().endswith("а") or parts[0].lower().endswith("я")):
        return "Ж"
    return "М"

def _parse_cadets_sheet(ws, default_group="", seen_groups=None):
    if seen_groups is None:
        seen_groups = []

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_idx = -1
    col_map = {}

    for r_i in range(min(15, len(rows))):
        row = [str(c or "").strip().lower() for c in rows[r_i]]
        has_fio = any(k in c for c in row for k in ['фио', 'ф.и.о.', 'фамили', 'курсант', 'слушател', 'обучающ'])
        has_rank = any(k in c for c in row for k in ['звание', 'в/зв', 'в/звание', 'воинск'])
        has_grp = any(k in c for c in row for k in ['групп', 'взвод', 'подразд', 'отделен', 'к/о', 'класс'])

        if (has_fio and (has_rank or has_grp)) or (has_fio and any('№' in c or 'п/п' in c for c in row)):
            header_idx = r_i
            for c_i, cell_text in enumerate(row):
                if any(k in cell_text for k in ['в/зв', 'звание', 'воинск']):
                    col_map['rank'] = c_i
                elif cell_text in ['фио', 'ф.и.о.', 'фамилия имя отчество', 'фамилия, имя, отчество', 'фамилия имя']:
                    col_map['fio'] = c_i
                elif 'сокр' in cell_text and 'фио' in cell_text:
                    col_map['short_fio'] = c_i
                elif any(k in cell_text for k in ['фио', 'ф.и.о.', 'фамилия']) and 'fio' not in col_map:
                    col_map['fio'] = c_i
                elif 'фамил' in cell_text and 'fio' not in col_map:
                    col_map['surname'] = c_i
                elif cell_text == 'имя' or cell_text.startswith('имя '):
                    col_map['name'] = c_i
                elif 'отчеств' in cell_text:
                    col_map['patronymic'] = c_i
                elif any(k in cell_text for k in ['групп', 'взвод', 'подразд', 'отделен', 'к/о', 'класс']):
                    col_map['group'] = c_i
                elif any(k in cell_text for k in ['пол', 'м/ж']):
                    col_map['gender'] = c_i
            break

    start_row = header_idx + 1 if header_idx != -1 else 0

    if 'fio' not in col_map and 'surname' not in col_map:
        col_scores = {}
        for r in rows[start_row:start_row + 15]:
            for c_i, cell in enumerate(r):
                val = str(cell or "").strip()
                if not val:
                    continue
                words = val.split()
                if len(words) in [2, 3] and any(ord('а') <= ord(ch.lower()) <= ord('я') for ch in val):
                    col_scores[c_i] = col_scores.get(c_i, 0) + 1
        if col_scores:
            best_fio_col = max(col_scores, key=col_scores.get)
            col_map['fio'] = best_fio_col
            if best_fio_col > 0:
                col_map['rank'] = best_fio_col - 1
            if best_fio_col + 1 < len(rows[start_row]):
                col_map['group'] = best_fio_col + 1
        else:
            col_map = {'rank': 0, 'fio': 1, 'group': 2}

    cadets = []
    current_group = default_group or "1"

    for r_i in range(start_row, len(rows)):
        row = rows[r_i]
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        fio = ""
        if 'fio' in col_map and col_map['fio'] < len(row) and row[col_map['fio']] is not None:
            fio = str(row[col_map['fio']]).strip()
        elif 'surname' in col_map:
            sur = str(row[col_map['surname']] or "").strip() if col_map['surname'] < len(row) else ""
            nm = str(row[col_map.get('name', 999)] or "").strip() if col_map.get('name', 999) < len(row) else ""
            pat = str(row[col_map.get('patronymic', 999)] or "").strip() if col_map.get('patronymic', 999) < len(row) else ""
            fio = " ".join(filter(None, [sur, nm, pat])).strip()

        fio_lower = fio.lower()
        if not fio or fio_lower in ['фио', 'ф.и.о.', 'фамилия', 'итого', 'всего']:
            continue

        if len(fio.split()) == 1 and ('взвод' in fio_lower or 'группа' in fio_lower or 'отделение' in fio_lower):
            current_group = fio
            continue

        rank = ""
        if 'rank' in col_map and col_map['rank'] < len(row) and row[col_map['rank']] is not None:
            rank = str(row[col_map['rank']]).strip()
            if rank.isdigit():
                rank = "курсант"
        if not rank:
            rank = "курсант"

        group_raw = None
        if 'group' in col_map and col_map['group'] < len(row):
            group_raw = row[col_map['group']]

        if group_raw is not None and str(group_raw).strip():
            if isinstance(group_raw, (int, float)):
                group_num = str(int(group_raw))
            else:
                group_num = str(group_raw).strip()
        else:
            group_num = current_group

        if group_num not in seen_groups:
            seen_groups.append(group_num)
        sort_order = seen_groups.index(group_num)

        short_fio = ""
        if 'short_fio' in col_map and col_map['short_fio'] < len(row) and row[col_map['short_fio']]:
            short_fio = str(row[col_map['short_fio']]).strip()
        if not short_fio:
            short_fio = make_short_fio(fio)

        raw_gender = None
        if 'gender' in col_map and col_map['gender'] < len(row):
            raw_gender = row[col_map['gender']]
        gender = detect_gender(fio, raw_gender)

        cadets.append({
            'rank': rank,
            'fio': fio,
            'short_fio': short_fio,
            'group_num': group_num,
            'gender': gender,
            'sort_order': sort_order
        })

    return cadets

def import_cadets_from_excel(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Файл не найден: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".xls":
        raise ValueError(
            "Файл имеет устаревший формат .xls (Excel 97-2003).\n"
            "Пожалуйста, откройте его в Microsoft Excel или LibreOffice и сохраните как «Книга Excel (.xlsx)»."
        )

    wb = openpyxl.load_workbook(file_path, data_only=True)

    priority_names = ['база', 'личный состав', 'курсанты', 'список', 'список личного состава', 'ведомость']
    target_sheets = []

    for sheet_name in wb.sheetnames:
        clean_name = sheet_name.strip().lower()
        if any(p in clean_name for p in priority_names):
            target_sheets.append(wb[sheet_name])
            break

    if not target_sheets:
        for sname in wb.sheetnames:
            ws = wb[sname]
            if ws.max_row and ws.max_row > 1:
                target_sheets.append(ws)

    if not target_sheets:
        target_sheets = [wb.active]

    all_cadets = []
    seen_groups = []

    for ws in target_sheets:
        default_grp = ws.title.strip() if len(target_sheets) > 1 else ""
        cadets = _parse_cadets_sheet(ws, default_group=default_grp, seen_groups=seen_groups)
        all_cadets.extend(cadets)

    return all_cadets

def style_range(ws, cell_range, border=None, font=None, alignment=None, fill=None):
    for row in ws[cell_range]:
        for cell in row:
            if border: cell.border = border
            if font: cell.font = font
            if alignment: cell.alignment = alignment
            if fill: cell.fill = fill

def export_report_to_excel(output_path, report_date, cadets, absences, unit_title, chief_title, chief_name, active_categories=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Расход"

    bold_font = Font(name="Times New Roman", size=12, bold=True)
    norm_font = Font(name="Times New Roman", size=12)
    title_font = Font(name="Times New Roman", size=12, bold=True)
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    thin = Side(border_style="thin", color="000000")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 1. Заголовки (строки 2-4)
    ws.merge_cells("B2:L2")
    ws["B2"] = "РАЗВЕРНУТАЯ СТРОЕВАЯ ЗАПИСКА"
    style_range(ws, "B2:L2", font=title_font, alignment=center)

    ws.merge_cells("B3:L3")
    ws["B3"] = unit_title
    style_range(ws, "B3:L3", font=norm_font, alignment=center)

    ws.merge_cells("B4:L4")
    try:
        dt_val = datetime.strptime(report_date, "%Y-%m-%d").strftime("%d.%m.%Y г.")
    except Exception:
        dt_val = report_date
    ws["B4"] = dt_val
    style_range(ws, "B4:L4", font=norm_font, alignment=center)

    ws.row_dimensions[4].height = 17.0

    # 2. Определение категорий для сводки
    standard_categories = ["Наряд", "Госпиталь", "Лазарет", "Командировка", "Похороны"]
    if active_categories:
        ordered_cats = [c for c in standard_categories if c in active_categories]
        other_cats = [c for c in active_categories if c not in standard_categories]
        categories = ordered_cats + other_cats
    else:
        categories = standard_categories

    group_order_map = {}
    for c in cadets:
        grp = c['group_num']
        if grp not in group_order_map:
            group_order_map[grp] = c.get('sort_order', 0)
    groups = sorted(group_order_map.keys(), key=lambda g: group_order_map[g])

    start_row = 5
    ws.row_dimensions[start_row].height = 18.0

    headers = ["№ п/п", "Подразд.", "П/С", "Н/Л"] + categories
    total_cols = len(headers) + 2
    end_col = 2 + total_cols - 1

    # Заголовки сводки
    for idx, h in enumerate(headers, start=2):
        cell = ws.cell(row=start_row, column=idx, value=h)
        cell.font = norm_font
        cell.fill = header_fill
        cell.border = cell_border
        cell.alignment = center

    # Колонка "ВСЕГО" (объединяется на 2 колонки K и L)
    v_col1 = 2 + len(headers)
    v_col2 = v_col1 + 1
    ws.merge_cells(start_row=start_row, start_column=v_col1, end_row=start_row, end_column=v_col2)
    ws.cell(row=start_row, column=v_col1, value="ВСЕГО")
    style_range(ws, f"{openpyxl.utils.get_column_letter(v_col1)}{start_row}:{openpyxl.utils.get_column_letter(v_col2)}{start_row}",
                border=cell_border, font=norm_font, fill=header_fill, alignment=center)

    curr_row = start_row + 1
    total_ps = 0
    total_nl = 0
    cat_totals = {cat: 0 for cat in categories}

    # Строки подразделений
    for idx, grp in enumerate(groups, 1):
        ws.row_dimensions[curr_row].height = 17.0
        grp_cadets = [c for c in cadets if c['group_num'] == grp]
        ps = len(grp_cadets)
        grp_abs = [a for a in absences if a['group_num'] == grp]

        row_counts = {}
        for cat in categories:
            cnt = sum(1 for a in grp_abs if a['category'] == cat)
            row_counts[cat] = cnt
            cat_totals[cat] += cnt

        grp_abs_total = len(grp_abs)
        nl = ps - grp_abs_total
        total_ps += ps
        total_nl += nl

        vals = [idx, grp, ps, nl] + [row_counts[cat] if row_counts[cat] > 0 else "" for cat in categories]
        for c_idx, val in enumerate(vals, start=2):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = norm_font
            cell.border = cell_border
            cell.alignment = center

        ws.merge_cells(start_row=curr_row, start_column=v_col1, end_row=curr_row, end_column=v_col2)
        ws.cell(row=curr_row, column=v_col1, value=grp_abs_total if grp_abs_total > 0 else "")
        style_range(ws, f"{openpyxl.utils.get_column_letter(v_col1)}{curr_row}:{openpyxl.utils.get_column_letter(v_col2)}{curr_row}",
                    border=cell_border, font=norm_font, alignment=center)
        curr_row += 1

    # Строки «Мужчины» и «Женщины»
    for g_label, g_val, num_idx in [("Мужчины", "М", len(groups) + 1), ("Женщины", "Ж", len(groups) + 2)]:
        ws.row_dimensions[curr_row].height = 17.0
        g_cadets = [c for c in cadets if c['gender'] == g_val]
        g_ps = len(g_cadets)
        g_abs = [a for a in absences if a['gender'] == g_val]
        g_nl = g_ps - len(g_abs)

        row_counts = {cat: sum(1 for a in g_abs if a['category'] == cat) for cat in categories}
        vals = [num_idx, g_label, g_ps, g_nl] + [row_counts[cat] if row_counts[cat] > 0 else "" for cat in categories]

        for c_idx, val in enumerate(vals, start=2):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = norm_font
            cell.border = cell_border
            cell.alignment = center

        ws.merge_cells(start_row=curr_row, start_column=v_col1, end_row=curr_row, end_column=v_col2)
        ws.cell(row=curr_row, column=v_col1, value=len(g_abs) if len(g_abs) > 0 else "")
        style_range(ws, f"{openpyxl.utils.get_column_letter(v_col1)}{curr_row}:{openpyxl.utils.get_column_letter(v_col2)}{curr_row}",
                    border=cell_border, font=norm_font, alignment=center)
        curr_row += 1

    # Строка «Всего»
    ws.row_dimensions[curr_row].height = 18.0
    total_abs_all = sum(cat_totals.values())
    summary_vals = [len(groups) + 3, "Всего", total_ps, total_nl] + [cat_totals[c] if cat_totals[c] > 0 else "" for c in categories]
    for c_idx, val in enumerate(summary_vals, start=2):
        cell = ws.cell(row=curr_row, column=c_idx, value=val)
        cell.font = bold_font
        cell.border = cell_border
        cell.alignment = center

    ws.merge_cells(start_row=curr_row, start_column=v_col1, end_row=curr_row, end_column=v_col2)
    ws.cell(row=curr_row, column=v_col1, value=total_abs_all if total_abs_all > 0 else "")
    style_range(ws, f"{openpyxl.utils.get_column_letter(v_col1)}{curr_row}:{openpyxl.utils.get_column_letter(v_col2)}{curr_row}",
                border=cell_border, font=bold_font, alignment=center)

    curr_row += 1
    ws.row_dimensions[curr_row].height = 14.0
    curr_row += 1

    # 3. Список отсутствующих
    last_col_letter = openpyxl.utils.get_column_letter(v_col2)
    ws.merge_cells(f"B{curr_row}:{last_col_letter}{curr_row}")
    ws[f"B{curr_row}"] = "Отсутствующие"
    style_range(ws, f"B{curr_row}:{last_col_letter}{curr_row}", font=norm_font, alignment=center)
    ws.row_dimensions[curr_row].height = 18.0
    curr_row += 1

    # Заголовки таблицы отсутствующих
    ws.row_dimensions[curr_row].height = 18.0
    ws.cell(row=curr_row, column=2, value="№ п/п").alignment = center
    ws.cell(row=curr_row, column=3, value="В/звание").alignment = center

    ws.merge_cells(start_row=curr_row, start_column=4, end_row=curr_row, end_column=6)
    ws.cell(row=curr_row, column=4, value="Фамилия и инициалы").alignment = center

    ws.merge_cells(start_row=curr_row, start_column=7, end_row=curr_row, end_column=8)
    ws.cell(row=curr_row, column=7, value="Причина отсутствия").alignment = center

    ws.cell(row=curr_row, column=9, value="Подразд.").alignment = center

    ws.merge_cells(start_row=curr_row, start_column=10, end_row=curr_row, end_column=v_col2)
    ws.cell(row=curr_row, column=10, value="Примечание").alignment = center

    style_range(ws, f"B{curr_row}:{last_col_letter}{curr_row}", border=cell_border, font=norm_font, fill=header_fill, alignment=center)
    curr_row += 1

    # Строки отсутствующих
    for idx, a in enumerate(absences, 1):
        ws.row_dimensions[curr_row].height = 18.0

        cell_num = ws.cell(row=curr_row, column=2, value=idx)
        cell_num.alignment = center
        cell_num.font = norm_font
        cell_num.border = cell_border

        cell_rank = ws.cell(row=curr_row, column=3, value=a['rank'])
        cell_rank.alignment = center
        cell_rank.font = norm_font
        cell_rank.border = cell_border

        ws.merge_cells(start_row=curr_row, start_column=4, end_row=curr_row, end_column=6)
        ws.cell(row=curr_row, column=4, value=a['short_fio'])
        style_range(ws, f"D{curr_row}:F{curr_row}", border=cell_border, font=norm_font, alignment=center)

        ws.merge_cells(start_row=curr_row, start_column=7, end_row=curr_row, end_column=8)
        ws.cell(row=curr_row, column=7, value=a['type_name'])
        style_range(ws, f"G{curr_row}:H{curr_row}", border=cell_border, font=norm_font, alignment=center)

        cell_grp = ws.cell(row=curr_row, column=9, value=a['group_num'])
        cell_grp.alignment = center
        cell_grp.font = norm_font
        cell_grp.border = cell_border

        ws.merge_cells(start_row=curr_row, start_column=10, end_row=curr_row, end_column=v_col2)
        note_text = str(a['note'] or "").strip()
        ws.cell(row=curr_row, column=10, value=note_text)
        note_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        style_range(ws, f"J{curr_row}:{last_col_letter}{curr_row}", border=cell_border, font=norm_font, alignment=note_align)

        curr_row += 1

    # 4. Подписи начальника
    curr_row += 1
    ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=v_col2)
    ws.cell(row=curr_row, column=2, value=chief_title.strip())
    style_range(ws, f"B{curr_row}:{last_col_letter}{curr_row}", font=norm_font, alignment=center)

    curr_row += 1
    ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=v_col2)
    ws.cell(row=curr_row, column=2, value=chief_name)
    style_range(ws, f"B{curr_row}:{last_col_letter}{curr_row}", font=norm_font, alignment=center)

    # 5. Ширина колонок
    col_widths = {
        'A': 3.0,
        'B': 7.0,
        'C': 11.5,
        'D': 8.5,
        'E': 9.0,
        'F': 13.0,
        'G': 12.5,
        'H': 11.5,
        'I': 15.0,
        'J': 15.0,
        'K': 9.0,
        'L': 13.0
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 6. Параметры страницы для идеальной печати на лист А4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3
    ws.page_margins.header = 0.0
    ws.page_margins.footer = 0.0

    ws.print_options.horizontalCentered = True
    ws.print_area = f"B2:{last_col_letter}{curr_row}"

    wb.save(output_path)
